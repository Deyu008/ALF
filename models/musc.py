import os
import sys
import gc
import numpy as np
import torch
import torch.nn.functional as F
sys.path.append('./models/backbone')

import datasets.mvtec as mvtec
from datasets.mvtec import _CLASSNAMES as _CLASSNAMES_mvtec_ad
import datasets.visa as visa
from datasets.visa import _CLASSNAMES as _CLASSNAMES_visa
import datasets.btad as btad
from datasets.btad import _CLASSNAMES as _CLASSNAMES_btad

import models.backbone.open_clip as open_clip
import models.backbone._backbones as _backbones
from models.modules._LNAMD import LNAMD
try:
    from models.modules._MSM import MSM, MSM_with_refs, MSM_with_bank
except ImportError:
    from models.modules._MSM import MSM
    MSM_with_refs = None
    MSM_with_bank = None
from models.modules._RsCIN import RsCIN
from models.modules.hier_fusion import HierarchicalFusion
from models.modules.adaptive_r_fusion import AdaptiveRFusion, AMSFConfig
from utils.metrics import compute_metrics
from openpyxl import Workbook
from tqdm import tqdm
import pickle
import time
import cv2

import warnings
warnings.filterwarnings("ignore")


class MuSc():
    def __init__(self, cfg, seed=0):
        self.cfg = cfg
        self.seed = seed
        self.device = torch.device("cuda:{}".format(cfg['device']) if torch.cuda.is_available() else "cpu")

        self.path = cfg['datasets']['data_path']
        self.dataset = cfg['datasets']['dataset_name']
        self.vis = cfg['testing']['vis']
        self.vis_type = cfg['testing']['vis_type']
        self.save_excel = cfg['testing']['save_excel']
        self.dump_rescore_inputs = bool(cfg.get('testing', {}).get('dump_rescore_inputs', False))
        self.dump_rescore_dir = str(cfg.get('testing', {}).get('dump_rescore_dir', 'rescore_inputs'))
        self.dump_rscin_hist = bool(cfg.get('testing', {}).get('dump_rscin_hist', False))
        self.dump_rscin_hist_dir = str(cfg.get('testing', {}).get('dump_rscin_hist_dir', 'rscin_hist'))
        # the categories to be tested
        self.categories = cfg['datasets']['class_name']
        if isinstance(self.categories, str):
            if self.categories.lower() == 'all':
                if self.dataset == 'visa':
                    self.categories = _CLASSNAMES_visa
                elif self.dataset == 'mvtec_ad':
                    self.categories = _CLASSNAMES_mvtec_ad
                elif self.dataset == 'btad':
                    self.categories = _CLASSNAMES_btad
            else:
                self.categories = [self.categories]

        self.model_name = cfg['models']['backbone_name']
        self.is_dinov3 = 'dinov3' in self.model_name.lower()
        self.image_size = cfg['datasets']['img_resize']
        self.batch_size = cfg['models']['batch_size']
        self.pretrained = cfg['models']['pretrained']
        self.features_list = [l+1 for l in cfg['models']['feature_layers']]
        self.divide_num = cfg['datasets']['divide_num']
        self.divide_iters = cfg['datasets'].get('divide_iters')
        self.r_list = cfg['models']['r_list']
        self.output_dir = os.path.join(cfg['testing']['output_dir'], self.dataset, self.model_name, 'imagesize{}'.format(self.image_size))
        os.makedirs(self.output_dir, exist_ok=True)
        self.load_backbone()
        self.autocast_enabled = torch.cuda.is_available() and (not self.is_dinov3)
        self.use_hier_fusion = bool(cfg['models'].get('use_hier_fusion', False))
        self.hier_fusion_cfg = cfg['models'].get('hier_fusion', {})
        self.hier_fusion = HierarchicalFusion(self.hier_fusion_cfg, self.device) if self.use_hier_fusion else None
        self.hier_fusion_stats = {} if self.use_hier_fusion else None

        # Adaptive R Fusion (AMSF) configuration
        self.adaptive_r_fusion_cfg = cfg['models'].get('adaptive_r_fusion', {})
        self.adaptive_r_fusion_config = AMSFConfig.from_dict(self.adaptive_r_fusion_cfg)
        self.adaptive_r_fusion = AdaptiveRFusion(self.adaptive_r_fusion_config, self.device) if self.adaptive_r_fusion_config.enabled else None
        self.adaptive_r_fusion_stats = {} if self.adaptive_r_fusion_config.enabled else None

        # Adaptive Layer Fusion (ALF): fuse MSM maps across backbone layers
        self.adaptive_layer_fusion_cfg = cfg['models'].get('adaptive_layer_fusion', {})
        self.adaptive_layer_fusion_config = AMSFConfig.from_dict(self.adaptive_layer_fusion_cfg)
        self.adaptive_layer_fusion = (
            AdaptiveRFusion(self.adaptive_layer_fusion_config, self.device)
            if self.adaptive_layer_fusion_config.enabled
            else None
        )
        self.adaptive_layer_fusion_stats = {} if self.adaptive_layer_fusion_config.enabled else None

        # Optional: L2-normalize image-level features (CLS / global embedding).
        # CLIP branch already normalizes by default; for DINO/DINOv3 this is off unless enabled.
        self.l2_normalize_image_features = bool(cfg.get('models', {}).get('l2_normalize_image_features', False))

        self.transductive_cfg = cfg.get('transductive', {})
        self.transductive_enabled = bool(self.transductive_cfg.get('enabled', False))
        if self.transductive_enabled and MSM_with_refs is None:
            warnings.warn("MSM_with_refs not available; disabling transductive refinement.")
            self.transductive_enabled = False
        self.transductive_num_passes = int(self.transductive_cfg.get('num_passes', 1))
        self.transductive_bank_ratio = float(self.transductive_cfg.get('bank_ratio', 0.4))
        # Bank selection strategy (worktree-compatible)
        self.transductive_bank_ratio_mode = str(self.transductive_cfg.get('bank_ratio_mode', 'fixed')).lower()
        self.transductive_bank_ratio_min = float(self.transductive_cfg.get('bank_ratio_min', 0.2))
        self.transductive_bank_ratio_max = float(self.transductive_cfg.get('bank_ratio_max', 0.7))
        self.transductive_bank_filter = bool(self.transductive_cfg.get('bank_filter', False))
        # Knee detection defaults (used when bank_ratio_mode=adaptive_knee)
        self.transductive_knee_strength_min = float(self.transductive_cfg.get('knee_strength_min', 0.03))
        self.transductive_knee_tail_frac = float(self.transductive_cfg.get('knee_tail_frac', 0.85))
        self.transductive_retrieval_topk = int(self.transductive_cfg.get('retrieval_topk', 0))
        self.transductive_score_pool = str(self.transductive_cfg.get('score_pool', 'q99')).lower()
        self.transductive_msm_topmin_min = float(self.transductive_cfg.get('msm_topmin_min', 0.0))
        self.transductive_msm_topmin_max = float(self.transductive_cfg.get('msm_topmin_max', 0.3))

        # Train-free transductive options (baseline-preserving; all default to old behavior)
        # - safe_fusion: blend Pass-0 and Pass-1 to reduce worst-case regressions
        # - apply_to_pixel=False (Route A): keep pixel maps from Pass-0, only change image-level score
        # - adaptive_gate: use a confidence-derived weight instead of fixed fusion weight
        self.transductive_safe_fusion = bool(self.transductive_cfg.get('safe_fusion', False))
        try:
            self.transductive_safe_fusion_weight = float(self.transductive_cfg.get('safe_fusion_weight', 0.5))
        except Exception:
            self.transductive_safe_fusion_weight = 0.5
        self.transductive_apply_to_pixel = bool(self.transductive_cfg.get('apply_to_pixel', True))
        self.transductive_adaptive_gate = bool(self.transductive_cfg.get('adaptive_gate', False))
        try:
            # Worktree default is 0.05; kept optional and only affects transductive when enabled.
            self.transductive_gate_threshold = float(self.transductive_cfg.get('gate_threshold', 0.05))
        except Exception:
            self.transductive_gate_threshold = 0.05


        # Inductive Normal Bank: use train-split normal images as MSM references.
        # This is train-free and inductive (no test-sample mutual scoring).
        self.inductive_bank_cfg = cfg.get('inductive_bank', {})
        self.inductive_bank_enabled = bool(self.inductive_bank_cfg.get('enabled', False))
        if self.inductive_bank_enabled and MSM_with_bank is None:
            warnings.warn("MSM_with_bank not available; disabling inductive normal bank.")
            self.inductive_bank_enabled = False
        self.inductive_bank_split = str(self.inductive_bank_cfg.get('split', 'train')).lower()
        try:
            self.inductive_bank_max_images = int(self.inductive_bank_cfg.get('max_images', -1))
        except Exception:
            self.inductive_bank_max_images = -1
        try:
            self.inductive_bank_query_chunk_size = int(self.inductive_bank_cfg.get('query_chunk_size', 1))
        except Exception:
            self.inductive_bank_query_chunk_size = 1
        self.inductive_bank_only_normal = bool(self.inductive_bank_cfg.get('only_normal', True))
        try:
            self.inductive_bank_msm_topmin_min = float(self.inductive_bank_cfg.get('msm_topmin_min', 0.0))
        except Exception:
            self.inductive_bank_msm_topmin_min = 0.0
        try:
            self.inductive_bank_msm_topmin_max = float(self.inductive_bank_cfg.get('msm_topmin_max', 0.3))
        except Exception:
            self.inductive_bank_msm_topmin_max = 0.3
        try:
            self.inductive_bank_retrieval_topk = float(self.inductive_bank_cfg.get('retrieval_topk', 0.0))
        except Exception:
            self.inductive_bank_retrieval_topk = 0.0
        self.inductive_bank_retrieval_topk = max(0.0, float(self.inductive_bank_retrieval_topk))
        self.inductive_bank_sample_mode = str(self.inductive_bank_cfg.get('sample_mode', 'sequential')).lower()
        try:
            self.inductive_bank_sample_seed = int(self.inductive_bank_cfg.get('sample_seed', self.seed))
        except Exception:
            self.inductive_bank_sample_seed = int(self.seed)

        if self.inductive_bank_enabled and self.transductive_enabled:
            warnings.warn(
                "inductive_bank.enabled=True is incompatible with transductive refinement; disabling transductive."
            )
            self.transductive_enabled = False

        self.image_pool = str(cfg.get('testing', {}).get('image_pool', 'max')).lower()

        # MSM peak-memory control: chunk cdist over reference images.
        try:
            self.msm_cdist_chunk_size = int(cfg.get('models', {}).get('msm_cdist_chunk_size', 0))
        except Exception:
            self.msm_cdist_chunk_size = 0


    def load_backbone(self):
        if 'dino' in self.model_name:
            # dino or dino_v2
            self.dino_model = _backbones.load(self.model_name)
            self.dino_model.to(self.device)
            self.preprocess = None
        else:
            # clip
            self.clip_model, _, self.preprocess = open_clip.create_model_and_transforms(self.model_name, self.image_size, pretrained=self.pretrained)
            self.clip_model.to(self.device)


    def load_datasets(self, category, divide_num=1, divide_iter=0):
        # dataloader
        if self.dataset == 'visa':
            test_dataset = visa.VisaDataset(source=self.path, split=visa.DatasetSplit.TEST,
                                            classname=category, resize=self.image_size, imagesize=self.image_size, clip_transformer=self.preprocess,
                                                divide_num=divide_num, divide_iter=divide_iter, random_seed=self.seed)
        elif self.dataset == 'mvtec_ad':
            test_dataset = mvtec.MVTecDataset(source=self.path, split=mvtec.DatasetSplit.TEST,
                                            classname=category, resize=self.image_size, imagesize=self.image_size, clip_transformer=self.preprocess,
                                                divide_num=divide_num, divide_iter=divide_iter, random_seed=self.seed)
        elif self.dataset == 'btad':
            test_dataset = btad.BTADDataset(source=self.path, split=btad.DatasetSplit.TEST,
                                            classname=category, resize=self.image_size, imagesize=self.image_size, clip_transformer=self.preprocess,
                                                divide_num=divide_num, divide_iter=divide_iter, random_seed=self.seed)
        return test_dataset


    def load_bank_datasets(self, category):
        """Load the dataset split used to build the inductive normal bank."""
        split = str(getattr(self, "inductive_bank_split", "train")).lower()
        if split in ("val", "valid", "validation"):
            split_mvtec = mvtec.DatasetSplit.VAL
            split_visa = visa.DatasetSplit.VAL
            split_btad = btad.DatasetSplit.VAL
        else:
            split_mvtec = mvtec.DatasetSplit.TRAIN
            split_visa = visa.DatasetSplit.TRAIN
            split_btad = btad.DatasetSplit.TRAIN

        if self.dataset == 'visa':
            return visa.VisaDataset(
                source=self.path,
                split=split_visa,
                classname=category,
                resize=self.image_size,
                imagesize=self.image_size,
                clip_transformer=self.preprocess,
                divide_num=1,
                divide_iter=0,
                random_seed=self.seed,
            )
        if self.dataset == 'mvtec_ad':
            return mvtec.MVTecDataset(
                source=self.path,
                split=split_mvtec,
                classname=category,
                resize=self.image_size,
                imagesize=self.image_size,
                clip_transformer=self.preprocess,
                divide_num=1,
                divide_iter=0,
                random_seed=self.seed,
            )
        if self.dataset == 'btad':
            return btad.BTADDataset(
                source=self.path,
                split=split_btad,
                classname=category,
                resize=self.image_size,
                imagesize=self.image_size,
                clip_transformer=self.preprocess,
                divide_num=1,
                divide_iter=0,
                random_seed=self.seed,
            )
        raise ValueError(f"Unknown dataset: {self.dataset}")


    def _build_inductive_normal_bank_r(self, category, r):
        """Build LNAMD features for train-split normal images (single r).

        Returns:
            bank_cache: Dict[str, torch.Tensor] where bank_cache[layer] has
                shape (N_bank, num_patches, C) on CPU (float16).
            None if bank is disabled or empty.
        """
        if not getattr(self, "inductive_bank_enabled", False):
            return None

        r_i = int(r)
        bank_dataset = self.load_bank_datasets(category)

        max_images = int(getattr(self, "inductive_bank_max_images", -1))
        if max_images == 0:
            return None

        target_n = len(bank_dataset)
        if max_images > 0:
            target_n = min(target_n, max_images)
        if target_n <= 0:
            return None

        sample_mode = str(getattr(self, "inductive_bank_sample_mode", "sequential")).lower()
        use_random_sampling = max_images > 0 and sample_mode in {"random", "shuffle"}
        loader_generator = None
        if use_random_sampling:
            seed_base = int(getattr(self, "inductive_bank_sample_seed", self.seed))
            category_seed = sum((idx + 1) * ord(ch) for idx, ch in enumerate(str(category)))
            sampled_seed = int((seed_base + 1315423911 + category_seed) % 2147483647)
            loader_generator = torch.Generator()
            loader_generator.manual_seed(sampled_seed)
            if r_i == int(self.r_list[0]):
                print(
                    f"[InductiveBank] {category}: sample_mode={sample_mode} max_images={max_images} sample_seed={sampled_seed}"
                )

        bank_loader = torch.utils.data.DataLoader(
            bank_dataset,
            batch_size=self.batch_size,
            shuffle=use_random_sampling,
            num_workers=0,
            pin_memory=True,
            generator=loader_generator,
        )

        lnamd_r = None
        bank_tensor = None
        bank_cls = None
        kept = 0

        for image_info in tqdm(bank_loader):
            if not isinstance(image_info, dict):
                continue
            image = image_info.get("image")
            if image is None:
                continue

            if getattr(self, "inductive_bank_only_normal", True):
                is_anomaly = image_info.get("is_anomaly")
                if is_anomaly is not None:
                    mask = (is_anomaly == 0)
                    if mask.numel() == 0 or not bool(mask.any().item()):
                        continue
                    image = image[mask]
                    if image.numel() == 0:
                        continue

            with torch.no_grad(), torch.cuda.amp.autocast(enabled=self.autocast_enabled):
                dtype = torch.float32 if self.is_dinov3 else torch.float
                input_image = image.to(dtype).to(self.device)

                if 'dinov2' in self.model_name:
                    patch_tokens = self.dino_model.get_intermediate_layers(
                        x=input_image,
                        n=[l - 1 for l in self.features_list],
                        return_class_token=False,
                    )
                    image_features = self.dino_model(input_image)
                    patch_tokens = [patch_tokens[l] for l in range(len(self.features_list))]
                    fake_cls = [torch.zeros_like(p)[:, 0:1, :] for p in patch_tokens]
                    patch_tokens = [torch.cat([fake_cls[i], patch_tokens[i]], dim=1) for i in range(len(patch_tokens))]
                elif self.is_dinov3:
                    patch_tokens_all = self.dino_model.get_intermediate_layers(
                        x=input_image, n=max(self.features_list)
                    )
                    image_features = self.dino_model(input_image)
                    patch_tokens = [patch_tokens_all[l - 1] for l in self.features_list]
                elif 'dino' in self.model_name:
                    patch_tokens_all = self.dino_model.get_intermediate_layers(
                        x=input_image, n=max(self.features_list)
                    )
                    image_features = self.dino_model(input_image)
                    patch_tokens = [patch_tokens_all[l - 1] for l in self.features_list]
                else:
                    image_features, patch_tokens = self.clip_model.encode_image(input_image, self.features_list)
                    image_features /= image_features.norm(dim=-1, keepdim=True)
                    patch_tokens = [patch_tokens[l] for l in range(len(self.features_list))]

                if self.l2_normalize_image_features and (self.is_dinov3 or ('dino' in self.model_name)):
                    image_features = image_features.float()
                    image_features = image_features / (image_features.norm(dim=-1, keepdim=True) + 1e-12)

                if self.use_hier_fusion and self.hier_fusion is not None:
                    saliency_output = self.hier_fusion.compute_batch(patch_tokens)
                    fused_attn = saliency_output["fused_attn"]
                    lambda_scale = saliency_output.get("lambda_scale")
                    patch_tokens = self.hier_fusion.apply_patch_gate(
                        patch_tokens, fused_attn, lambda_scale=lambda_scale
                    )

            if lnamd_r is None:
                feature_dim = int(patch_tokens[0].shape[-1])
                lnamd_r = LNAMD(
                    device=self.device,
                    r=r_i,
                    feature_dim=feature_dim,
                    feature_layer=self.features_list,
                )

            image_features_cpu = image_features.detach().to(dtype=torch.float32).cpu()

            with torch.no_grad(), torch.cuda.amp.autocast(enabled=self.autocast_enabled):
                features = lnamd_r._embed(patch_tokens)

            # LNAMD._embed returns CPU tensors; normalize in float32 then store in float16.
            features = features.to(dtype=torch.float32)
            features = features / (features.norm(dim=-1, keepdim=True) + 1e-12)
            features = features.to(dtype=torch.float16)

            if bank_tensor is None:
                _, patch_num, num_layers, c = features.shape
                bank_tensor = torch.empty(
                    (target_n, patch_num, num_layers, c),
                    dtype=torch.float16,
                    device="cpu",
                )
                bank_cls = torch.empty(
                    (target_n, int(image_features_cpu.shape[-1])),
                    dtype=torch.float16,
                    device="cpu",
                )

            remaining = target_n - kept
            if remaining <= 0:
                break
            b = int(features.shape[0])
            if b > remaining:
                features = features[:remaining]
                image_features_cpu = image_features_cpu[:remaining]
                b = remaining

            bank_tensor[kept : kept + b].copy_(features)
            if bank_cls is not None:
                bank_cls[kept : kept + b].copy_(image_features_cpu[:b].to(dtype=torch.float16))
            kept += b
            if kept >= target_n:
                break

        if kept <= 0 or bank_tensor is None:
            return None

        bank_tensor = bank_tensor[:kept]
        bank_cache = {str(li): bank_tensor[:, :, li, :] for li in range(bank_tensor.shape[2])}
        if bank_cls is not None:
            bank_cache['cls'] = bank_cls[:kept]
        print(f"[InductiveBank] {category} r={r_i}: split={self.inductive_bank_split} bank_n={kept}")
        return bank_cache


    def visualization(self, image_path_list, gt_list, pr_px, category):
        def normalization01(img):
            return (img - img.min()) / (img.max() - img.min())
        if self.vis_type == 'single_norm':
            # normalized per image
            for i, path in enumerate(image_path_list):
                anomaly_type = path.split('/')[-2]
                img_name = path.split('/')[-1]
                if anomaly_type not in ['good', 'Normal', 'ok'] and gt_list[i] != 0:
                    save_path = os.path.join(self.output_dir, category, anomaly_type)
                    os.makedirs(save_path, exist_ok=True)
                    save_path = os.path.join(save_path, img_name)
                    anomaly_map = pr_px[i].squeeze()
                    anomaly_map = normalization01(anomaly_map)*255
                    anomaly_map = cv2.applyColorMap(anomaly_map.astype(np.uint8), cv2.COLORMAP_JET)
                    cv2.imwrite(save_path, anomaly_map)
        else:
            # normalized all image
            pr_px = normalization01(pr_px)
            for i, path in enumerate(image_path_list):
                anomaly_type = path.split('/')[-2]
                img_name = path.split('/')[-1]
                save_path = os.path.join(self.output_dir, category, anomaly_type)
                os.makedirs(save_path, exist_ok=True)
                save_path = os.path.join(save_path, img_name)
                anomaly_map = pr_px[i].squeeze()
                anomaly_map *= 255
                anomaly_map = cv2.applyColorMap(anomaly_map.astype(np.uint8), cv2.COLORMAP_JET)
                cv2.imwrite(save_path, anomaly_map)


    def _fuse_anomaly_maps_r(self, anomaly_maps_r: torch.Tensor, category: str) -> torch.Tensor:
        """Fuse anomaly maps from different r values.
        
        Args:
            anomaly_maps_r: (num_r, N, num_patches) anomaly maps from different r values
            category: current category name (for logging)
            
        Returns:
            Fused anomaly maps: (N, num_patches)
        """
        if self.adaptive_r_fusion is not None and self.adaptive_r_fusion_config.enabled:
            # AMSF fusion (optionally calibrated to equal-weight baseline).
            baseline_map = anomaly_maps_r.mean(dim=0)
            result = self.adaptive_r_fusion.fuse(anomaly_maps_r, baseline_map=baseline_map)
            
            # Log weight statistics for this category
            if self.adaptive_r_fusion_config.log_weights:
                weight_stats = self.adaptive_r_fusion.get_weight_stats()
                if weight_stats:
                    msg = (f"[AMSF] {category}: mean_weights={weight_stats['mean_weights']}, "
                           f"std={weight_stats['std_weights']}")
                    if "mean_gate" in weight_stats:
                        msg += (f", mean_gate={weight_stats['mean_gate']:.4f}, "
                                f"std_gate={weight_stats['std_gate']:.4f}, "
                                f"mean_consistency={weight_stats['mean_consistency']:.4f}, "
                                f"method={weight_stats['consistency_method']}, "
                                f"tau={weight_stats['consistency_tau']}, "
                                f"beta={weight_stats['consistency_beta']}, "
                                f"topk_ratio={weight_stats['consistency_topk_ratio']}, "
                                f"hard_floor={weight_stats.get('consistency_gate_hard_floor', 0.0)}")
                    if "mean_patch_gate" in weight_stats:
                        msg += (f", patch_gate={weight_stats['mean_patch_gate']:.4f}±{weight_stats['std_patch_gate']:.4f}, "
                                f"pg_method={weight_stats.get('patch_gate_method', '')}, "
                                f"pg_tau={weight_stats.get('patch_gate_tau', 0.0)}, "
                                f"pg_beta={weight_stats.get('patch_gate_beta', 0.0)}, "
                                f"pg_k={weight_stats.get('patch_gate_smooth_kernel', 0)}, "
                                f"pg_floor={weight_stats.get('patch_gate_floor', 0.0)}, "
                                f"pg_max={weight_stats.get('patch_gate_max', 1.0)}, "
                                f"pg_protect_q={weight_stats.get('patch_gate_protect_quantile', -1.0)}, "
                                f"pg_protect_src={weight_stats.get('patch_gate_protect_source', 'reference')}, "
                                f"pg_power={weight_stats.get('patch_gate_power', 1.0)}, "
                                f"pg_mode={weight_stats.get('patch_gate_mode', 'all')}, "
                                f"pg_bg_p={weight_stats.get('patch_gate_bg_percentile', 0.9)}, "
                                f"pg_cons_th={weight_stats.get('patch_gate_consistency_threshold', -1.0)}, "
                                f"pg_gate_th={weight_stats.get('patch_gate_gate_threshold', -1.0)}, "
                                f"pg_apply={weight_stats.get('patch_gate_apply_logic', 'or')}")
                    # Optional (may be absent depending on fusion module version)
                    if "mean_residual_mask" in weight_stats:
                        msg += (f", mean_rmask={weight_stats['mean_residual_mask']:.4f}, "
                                f"std_rmask={weight_stats['std_residual_mask']:.4f}, "
                                f"rmask_p={weight_stats.get('residual_mask_bg_percentile', 0.0):.3f}, "
                                f"rmask_temp={weight_stats.get('residual_mask_temperature', 0.0):.4f}, "
                                f"rmask_floor={weight_stats.get('residual_mask_floor', 0.0):.3f}, "
                                f"rmask_power={weight_stats.get('residual_mask_power', 1.0):.3f}, "
                                f"rmask_pos_only={weight_stats.get('residual_mask_pos_only', True)}")
                    if "mean_calib_scale" in weight_stats:
                        msg += (f", calib={weight_stats.get('calibration_method', '')}, "
                                f"scale={weight_stats['mean_calib_scale']:.4f}±{weight_stats['std_calib_scale']:.4f}, "
                                f"shift={weight_stats['mean_calib_shift']:.4f}±{weight_stats['std_calib_shift']:.4f}")
                        msg += f", calib_n={weight_stats.get('num_calib_samples', 0)}"
                        if "calibration_bg_percentile" in weight_stats:
                            msg += (f", bg_p={weight_stats['calibration_bg_percentile']:.3f}, "
                                    f"bg_n={weight_stats.get('mean_bg_count', 0.0):.1f}")
                        if "calibration_tail_q0" in weight_stats:
                            msg += (f", tail_q0={weight_stats['calibration_tail_q0']:.3f}, "
                                    f"tail_q1={weight_stats['calibration_tail_q1']:.3f}, "
                                    f"tail_scale=[{weight_stats.get('calibration_tail_scale_min', 0.0):.2f},"
                                    f"{weight_stats.get('calibration_tail_scale_max', 0.0):.2f}], "
                                    f"tail_clamp={weight_stats.get('calibration_tail_clamp_to_ref', False)}")
                    print(msg)
                    if self.adaptive_r_fusion_stats is not None:
                        self.adaptive_r_fusion_stats[category] = weight_stats
                # Reset for next category
                self.adaptive_r_fusion.reset_stats()
            
            return result
        else:
            # Default: equal-weight averaging
            return torch.mean(anomaly_maps_r, dim=0)


    def _fuse_anomaly_maps_layers(self, anomaly_maps_layers: torch.Tensor, category: str, r: int) -> torch.Tensor:
        """Fuse anomaly maps across backbone layers for a given r.

        Args:
            anomaly_maps_layers: (num_layers, N, num_patches) per-layer MSM anomaly maps
            category: current category name (for logging)
            r: current aggregation degree (for logging)

        Returns:
            Fused anomaly maps: (N, num_patches)
        """
        if self.adaptive_layer_fusion is not None and self.adaptive_layer_fusion_config.enabled:
            baseline_map = anomaly_maps_layers.mean(dim=0)
            result = self.adaptive_layer_fusion.fuse(anomaly_maps_layers, baseline_map=baseline_map)

            if self.adaptive_layer_fusion_config.log_weights:
                weight_stats = self.adaptive_layer_fusion.get_weight_stats()
                if weight_stats:
                    msg = (f"[ALF] {category} r={r}: mean_weights={weight_stats['mean_weights']}, "
                           f"std={weight_stats['std_weights']}")
                    if "mean_gate" in weight_stats:
                        msg += (f", mean_gate={weight_stats['mean_gate']:.4f}, "
                                f"std_gate={weight_stats['std_gate']:.4f}, "
                                f"mean_consistency={weight_stats['mean_consistency']:.4f}, "
                                f"method={weight_stats['consistency_method']}, "
                                f"tau={weight_stats['consistency_tau']}, "
                                f"beta={weight_stats['consistency_beta']}, "
                                f"topk_ratio={weight_stats['consistency_topk_ratio']}, "
                                f"hard_floor={weight_stats.get('consistency_gate_hard_floor', 0.0)}")
                    if "mean_calib_scale" in weight_stats:
                        msg += (f", calib={weight_stats.get('calibration_method', '')}, "
                                f"scale={weight_stats['mean_calib_scale']:.4f}±{weight_stats['std_calib_scale']:.4f}, "
                                f"shift={weight_stats['mean_calib_shift']:.4f}±{weight_stats['std_calib_shift']:.4f}")
                        msg += f", calib_n={weight_stats.get('num_calib_samples', 0)}"
                        if "calibration_bg_percentile" in weight_stats:
                            msg += (f", bg_p={weight_stats['calibration_bg_percentile']:.3f}, "
                                    f"bg_n={weight_stats.get('mean_bg_count', 0.0):.1f}")
                    print(msg)
                    if self.adaptive_layer_fusion_stats is not None:
                        self.adaptive_layer_fusion_stats[f"{category}_r{r}"] = weight_stats
                self.adaptive_layer_fusion.reset_stats()

            return result

        return torch.mean(anomaly_maps_layers, dim=0)



    def make_category_data(self, category):
        print(category)

        if self.use_hier_fusion and self.hier_fusion is not None:
            self.hier_fusion.reset_stats()

        # divide sub-datasets
        divide_num = self.divide_num
        need_transductive = self.transductive_enabled and self.transductive_num_passes > 1
        anomaly_maps = torch.tensor([]).double()
        anomaly_maps_base = torch.tensor([]).double() if need_transductive else None
        gt_list = []
        img_masks = []
        class_tokens = []
        cls_token_batches = []
        image_path_list = []
        start_time_all = time.time()
        dataset_num = 0

        divide_iters = self.divide_iters
        if divide_iters is None:
            divide_iters = list(range(divide_num))
        elif isinstance(divide_iters, int):
            divide_iters = [divide_iters]
        else:
            divide_iters = list(divide_iters)
        divide_iters = sorted(set(int(i) for i in divide_iters))
        for i in divide_iters:
            if i < 0 or i >= divide_num:
                raise ValueError(f"Invalid divide_iter={i} for divide_num={divide_num}.")
        if need_transductive and len(divide_iters) != 1:
            raise ValueError(
                "当前实现的 transductive refinement 需要单个 divide_iter（保证 bank/ref 索引一致）。"
                "请设置 `datasets.divide_num=1`，或设置 `datasets.divide_iters` 为单元素列表（如 [0]）。"
            )

        for divide_iter in divide_iters:
            test_dataset = self.load_datasets(category, divide_num=divide_num, divide_iter=divide_iter)
            test_dataloader = torch.utils.data.DataLoader(
                test_dataset,
                batch_size=self.batch_size,
                shuffle=False,
                num_workers=0,
                pin_memory=True,
            )

            # extract features
            patch_tokens_list = []
            cls_token_batches_iter = []
            saliency_batches = [] if self.use_hier_fusion and self.hier_fusion is not None else None
            lambda_batches = [] if self.use_hier_fusion and self.hier_fusion is not None else None
            subset_num = len(test_dataset)
            dataset_num += subset_num
            start_time = time.time()
            for image_info in tqdm(test_dataloader):
            # for image_info in test_dataloader:
                if isinstance(image_info, dict):
                    image = image_info["image"]
                    image_path_list.extend(image_info["image_path"])
                    img_masks.append(image_info["mask"])
                    gt_list.extend(list(image_info["is_anomaly"].numpy()))
                with torch.no_grad(), torch.cuda.amp.autocast(enabled=self.autocast_enabled):
                    dtype = torch.float32 if self.is_dinov3 else torch.float
                    input_image = image.to(dtype).to(self.device)
                    if 'dinov2' in self.model_name:
                        patch_tokens = self.dino_model.get_intermediate_layers(x=input_image, n=[l-1 for l in self.features_list], return_class_token=False)
                        image_features = self.dino_model(input_image)
                        patch_tokens = [patch_tokens[l] for l in range(len(self.features_list))]
                        fake_cls = [torch.zeros_like(p)[:, 0:1, :] for p in patch_tokens]
                        patch_tokens = [torch.cat([fake_cls[i], patch_tokens[i]], dim=1) for i in range(len(patch_tokens))]
                    elif self.is_dinov3:
                        patch_tokens_all = self.dino_model.get_intermediate_layers(x=input_image, n=max(self.features_list))
                        image_features = self.dino_model(input_image)
                        patch_tokens = [patch_tokens_all[l-1] for l in self.features_list]
                    elif 'dino' in self.model_name:
                        patch_tokens_all = self.dino_model.get_intermediate_layers(x=input_image, n=max(self.features_list))
                        image_features = self.dino_model(input_image)
                        patch_tokens = [patch_tokens_all[l-1] for l in self.features_list]
                    else: # clip
                        image_features, patch_tokens = self.clip_model.encode_image(input_image, self.features_list)
                        image_features /= image_features.norm(dim=-1, keepdim=True)
                        patch_tokens = [patch_tokens[l] for l in range(len(self.features_list))]

                    # Keep feature scale consistent across backbones when desired.
                    # (trainfree branch normalized DINO features; mainline keeps it optional for baseline safety)
                    if self.l2_normalize_image_features and (self.is_dinov3 or ('dino' in self.model_name)):
                        image_features = image_features.float()
                        image_features = image_features / (image_features.norm(dim=-1, keepdim=True) + 1e-12)
                image_features_cpu = image_features.detach().cpu()
                cls_token_batches.append(image_features_cpu)
                cls_token_batches_iter.append(image_features_cpu)
                image_features_np = [image_features_cpu[bi].squeeze().numpy() for bi in range(image_features_cpu.shape[0])]
                class_tokens.extend(image_features_np)

                if self.use_hier_fusion and self.hier_fusion is not None:
                    saliency_output = self.hier_fusion.compute_batch(patch_tokens)
                    fused_attn = saliency_output["fused_attn"]
                    lambda_scale = saliency_output.get("lambda_scale")
                    patch_tokens = self.hier_fusion.apply_patch_gate(patch_tokens, fused_attn, lambda_scale=lambda_scale)
                    saliency_batches.append(fused_attn.cpu())
                    if lambda_batches is not None:
                        if lambda_scale is not None:
                            lambda_batches.append(lambda_scale.detach().cpu())
                        else:
                            lambda_batches.append(None)

                patch_tokens_list.append([p.cpu() for p in patch_tokens])  # (B, L+1, C)
            end_time = time.time()
            print('extract time: {}ms per image'.format((end_time-start_time)*1000/subset_num))

            cls_tensor_iter = None
            if cls_token_batches_iter:
                cls_tensor_iter = torch.cat(cls_token_batches_iter, dim=0).float()
                cls_tensor_iter = cls_tensor_iter / (cls_tensor_iter.norm(dim=-1, keepdim=True) + 1e-12)

            # LNAMD
            feature_dim = patch_tokens_list[0][0].shape[-1]
            fused_attn_all = None
            lambda_all = None
            if self.use_hier_fusion and self.hier_fusion is not None and saliency_batches is not None:
                fused_attn_all = torch.cat(saliency_batches, dim=0).to(self.device)
                if lambda_batches is not None and any(lb is not None for lb in lambda_batches):
                    lambda_tensors = [lb for lb in lambda_batches if lb is not None]
                    if lambda_tensors:
                        lambda_all = torch.cat(lambda_tensors, dim=0).to(self.device)

            Z_cache = {} if need_transductive else None

            # Collect anomaly maps for each r value
            anomaly_maps_r_list = []
            for r in self.r_list:
                start_time = time.time()
                print('aggregation degree: {}'.format(r))

                bank_cache_r = None
                if self.inductive_bank_enabled:
                    bank_cache_r = self._build_inductive_normal_bank_r(category, int(r))
                    if bank_cache_r is None:
                        print(f"[InductiveBank] {category} r={r}: empty bank -> fallback to mutual scoring")

                bank_retrieval_indices = None
                if (
                    bank_cache_r is not None
                    and float(getattr(self, 'inductive_bank_retrieval_topk', 0.0)) > 0.0
                    and isinstance(bank_cache_r, dict)
                    and 'cls' in bank_cache_r
                    and cls_tensor_iter is not None
                ):
                    bank_cls = bank_cache_r.get('cls')
                    if isinstance(bank_cls, torch.Tensor) and bank_cls.numel() > 0:
                        bank_cls_f = bank_cls.float()
                        bank_cls_f = bank_cls_f / (bank_cls_f.norm(dim=-1, keepdim=True) + 1e-12)
                        bank_n = int(bank_cls_f.shape[0])
                        k_cfg = float(getattr(self, 'inductive_bank_retrieval_topk', 0.0))
                        if k_cfg >= 1.0:
                            k = int(k_cfg)
                        else:
                            k = int(round(bank_n * k_cfg))
                        if k_cfg > 0.0 and k <= 0:
                            k = 1
                        k = max(0, min(int(k), int(bank_n)))
                        if 0 < k < bank_n:
                            sim = cls_tensor_iter @ bank_cls_f.T
                            bank_retrieval_indices = torch.topk(sim, k=k, dim=1, largest=True).indices
                            print(f"[InductiveBank] {category} r={r}: retrieval_topk={k}/{bank_n}")

                LNAMD_r = LNAMD(device=self.device, r=r, feature_dim=feature_dim, feature_layer=self.features_list)
                Z_layers = {}
                for im in range(len(patch_tokens_list)):
                    patch_tokens = [p.to(self.device) for p in patch_tokens_list[im]]
                    with torch.no_grad(), torch.cuda.amp.autocast(enabled=self.autocast_enabled):
                        features = LNAMD_r._embed(patch_tokens)
                        features /= features.norm(dim=-1, keepdim=True)
                        for l in range(len(self.features_list)):
                            if str(l) not in Z_layers.keys():
                                Z_layers[str(l)] = []
                            Z_layers[str(l)].append(features[:, :, l, :])
                end_time = time.time()
                print('LNAMD-{}: {}ms per image'.format(r, (end_time-start_time)*1000/subset_num))

                # MSM
                anomaly_maps_l = torch.tensor([]).double()
                start_time = time.time()
                if need_transductive and Z_cache is not None:
                    Z_cache[int(r)] = {}
                for l in Z_layers.keys():
                    # different layers
                    Z_cpu = torch.cat(Z_layers[l], dim=0)  # (N, L, C), on CPU
                    if need_transductive and Z_cache is not None:
                        Z_cache[int(r)][l] = Z_cpu
                    Z = Z_cpu.to(self.device)  # (N, L, C)
                    print('layer-{} mutual scoring...'.format(l))

                    if bank_cache_r is not None and MSM_with_bank is not None and str(l) in bank_cache_r:
                        anomaly_maps_msm = MSM_with_bank(
                            Z_query=Z,
                            device=self.device,
                            Z_bank=bank_cache_r[str(l)],
                            topmin_min=self.inductive_bank_msm_topmin_min,
                            topmin_max=self.inductive_bank_msm_topmin_max,
                            bank_indices=bank_retrieval_indices,
                            query_chunk_size=self.inductive_bank_query_chunk_size,
                            cdist_chunk_size=self.msm_cdist_chunk_size,
                        )
                    else:
                        anomaly_maps_msm = MSM(
                            Z=Z,
                            device=self.device,
                            topmin_min=0,
                            topmin_max=0.3,
                            cdist_chunk_size=self.msm_cdist_chunk_size,
                        )
                    if fused_attn_all is not None:
                        anomaly_maps_msm = self.hier_fusion.rescore_patch_scores(
                            anomaly_maps_msm,
                            fused_attn_all,
                            lambda_scale=lambda_all,
                        )
                    anomaly_maps_l = torch.cat((anomaly_maps_l, anomaly_maps_msm.unsqueeze(0).cpu()), dim=0)
                    torch.cuda.empty_cache()
                anomaly_maps_l = self._fuse_anomaly_maps_layers(anomaly_maps_l, category, int(r))
                anomaly_maps_r_list.append(anomaly_maps_l)
                del bank_cache_r
                gc.collect()
                end_time = time.time()
                print('MSM: {}ms per image'.format((end_time-start_time)*1000/subset_num))
            
            # Stack and fuse anomaly maps from different r values
            anomaly_maps_r = torch.stack(anomaly_maps_r_list, dim=0)  # (num_r, N, num_patches)
            anomaly_maps_patch_base = self._fuse_anomaly_maps_r(anomaly_maps_r, category).to(self.device)
            del anomaly_maps_r, anomaly_maps_r_list
            torch.cuda.empty_cache()

            def pool_scores(scores_2d: torch.Tensor, mode: str) -> np.ndarray:
                scores_np = scores_2d.detach().float().cpu().numpy()
                if scores_np.ndim != 2:
                    raise ValueError(f"Expected 2D scores (N, P), got {scores_np.shape}.")
                if mode == 'max':
                    return scores_np.max(axis=1)
                if mode in ('q99', 'p99'):
                    return np.quantile(scores_np, 0.99, axis=1)
                if mode == 'topk_mean':
                    k = max(1, int(scores_np.shape[1] * 0.01))
                    topk = np.partition(scores_np, -k, axis=1)[:, -k:]
                    return topk.mean(axis=1)
                raise ValueError(f"Unknown score pool mode: {mode}")

            anomaly_maps_patch = anomaly_maps_patch_base
            bank_indices = None
            transductive_knee_strength = 0.0
            transductive_kept_ratio = 1.0
            if need_transductive and Z_cache is not None:
                n_images = int(anomaly_maps_patch_base.shape[0])
                if n_images >= 2:
                    score_img0 = pool_scores(anomaly_maps_patch_base, self.transductive_score_pool)

                    # --- Bank selection (worktree-compatible) ---
                    ratio_min = float(self.transductive_bank_ratio_min)
                    ratio_max = float(self.transductive_bank_ratio_max)
                    ratio_min = max(0.01, min(0.99, ratio_min))
                    ratio_max = max(0.01, min(0.99, ratio_max))
                    if ratio_max < ratio_min:
                        ratio_min, ratio_max = ratio_max, ratio_min

                    min_size = max(2, int(np.ceil(n_images * ratio_min)))
                    max_size = min(n_images - 1, int(np.floor(n_images * ratio_max)))
                    if max_size < min_size:
                        max_size = min_size

                    fallback_ratio = float(self.transductive_bank_ratio)
                    fallback_ratio = max(0.05, min(0.95, fallback_ratio))
                    fallback_size = int(np.floor(n_images * fallback_ratio))
                    fallback_size = max(min_size, min(max_size, fallback_size))

                    sorted_idx = np.argsort(score_img0).astype(np.int64)
                    sorted_scores = score_img0[sorted_idx]

                    mode = str(self.transductive_bank_ratio_mode).lower()
                    bank_size_raw = None

                    if mode == 'fixed':
                        bank_size = fallback_size
                    elif mode in ('adaptive_knee', 'knee'):
                        bank_size = fallback_size
                        if n_images >= 20 and float(np.ptp(sorted_scores)) > 1e-8:
                            y = (sorted_scores - sorted_scores.min()) / (sorted_scores.max() - sorted_scores.min() + 1e-12)
                            win = 5
                            if win > 1 and y.shape[0] >= win:
                                kernel = np.ones(win, dtype=np.float64) / float(win)
                                pad = win // 2
                                y_pad = np.pad(y, (pad, pad), mode='edge')
                                y = np.convolve(y_pad, kernel, mode='valid')
                            x = np.linspace(0.0, 1.0, num=y.shape[0], dtype=np.float64)
                            d = (y.astype(np.float64) - x.astype(np.float64))
                            if d.size > 0 and np.all(np.isfinite(d)):
                                knee_strength_min = float(self.transductive_knee_strength_min)
                                knee_tail_frac = float(self.transductive_knee_tail_frac)
                                knee_tail_frac = max(0.0, min(1.0, knee_tail_frac))

                                if float(np.mean(d)) < 0.0:
                                    d_use = -d
                                    knee_shape = 'convex'
                                else:
                                    d_use = d
                                    knee_shape = 'concave'

                                knee_idx = int(np.argmax(d_use))
                                knee_strength = float(abs(d[knee_idx]))
                                knee_rel = float(knee_idx) / float(max(1, int(y.shape[0]) - 1))

                                bank_size_raw = int(knee_idx + 1)
                                bank_size_knee = max(min_size, min(max_size, bank_size_raw))

                                knee_is_clamped = (bank_size_knee >= int(max_size)) and (bank_size_raw > int(max_size))
                                knee_is_tail = knee_rel >= knee_tail_frac
                                knee_is_weak = knee_strength < knee_strength_min

                                if (knee_is_tail or knee_is_weak or knee_is_clamped) and fallback_size != bank_size_knee:
                                    bank_size = fallback_size
                                    transductive_knee_strength = knee_strength * 0.5
                                    print(
                                        f"[Transductive] adaptive_knee fallback: shape={knee_shape} "
                                        f"knee_idx={knee_idx} rel={knee_rel:.3f} strength={knee_strength:.4f} "
                                        f"(tail>={knee_tail_frac:.2f} or strength<{knee_strength_min:.4f} or clamped)"
                                    )
                                else:
                                    bank_size = bank_size_knee
                                    transductive_knee_strength = knee_strength
                                    print(
                                        f"[Transductive] adaptive_knee: shape={knee_shape} "
                                        f"knee_idx={knee_idx} rel={knee_rel:.3f} strength={knee_strength:.4f}"
                                    )
                    else:
                        raise ValueError(f"Unknown transductive.bank_ratio_mode: {mode}")

                    chosen_ratio = float(bank_size) / float(n_images)
                    bank_indices = sorted_idx[:bank_size].astype(np.int64)
                    print(
                        f"[Transductive] bank_ratio={chosen_ratio:.3f} bank_size={bank_size}/{n_images} pool={self.transductive_score_pool}"
                    )
                    if mode != 'fixed':
                        print(
                            f"[Transductive] bank_ratio_mode={mode} raw_bank_size={bank_size_raw} "
                            f"fallback_size={fallback_size} clamp=[{min_size},{max_size}]"
                        )

                    bank_tensor = torch.as_tensor(bank_indices, dtype=torch.long)

                    # --- Bank decontamination (knn_mad) ---
                    if (
                        self.transductive_bank_filter
                        and bank_tensor.numel() >= 10
                        and cls_token_batches
                    ):
                        cls_tensor = torch.cat(cls_token_batches, dim=0).float()
                        cls_tensor = cls_tensor / (cls_tensor.norm(dim=-1, keepdim=True) + 1e-12)
                        bank_cls = cls_tensor.index_select(0, bank_tensor)
                        sim_bank = bank_cls @ bank_cls.T
                        sim_bank.fill_diagonal_(-1e9)

                        k_density = min(5, int(bank_tensor.numel() - 1))
                        if k_density > 0:
                            density = torch.topk(sim_bank, k=k_density, dim=1, largest=True).values.mean(dim=1)
                            med = density.median()
                            mad = (density - med).abs().median() + 1e-12
                            thr = med - 2.0 * mad
                            keep_mask = density >= thr
                            keep_num = int(keep_mask.sum().item())

                            if keep_num < int(min_size):
                                top_idx = torch.topk(density, k=int(min_size), dim=0, largest=True).indices
                                keep_mask = torch.zeros_like(keep_mask, dtype=torch.bool)
                                keep_mask[top_idx] = True
                                keep_num = int(min_size)

                            if keep_num < int(bank_tensor.numel()):
                                old_num = int(bank_tensor.numel())
                                bank_tensor = bank_tensor[keep_mask]
                                bank_indices = bank_tensor.detach().cpu().numpy().astype(np.int64)
                                transductive_kept_ratio = float(keep_num) / float(old_num)
                                print(
                                    f"[Transductive] bank_filter=knn_mad kept={keep_num}/{old_num} "
                                    f"k={k_density} thr={float(thr):.4f}"
                                )

                    ref_indices = bank_tensor
                    topk = int(self.transductive_retrieval_topk)
                    if topk > 0 and bank_tensor.numel() > 1 and cls_token_batches:
                        cls_tensor = torch.cat(cls_token_batches, dim=0).float()
                        cls_tensor = cls_tensor / (cls_tensor.norm(dim=-1, keepdim=True) + 1e-12)
                        bank_cls = cls_tensor.index_select(0, bank_tensor)
                        sim = cls_tensor @ bank_cls.T  # (N, R)
                        # exclude self if it is inside bank
                        pos_in_bank = torch.full((cls_tensor.shape[0],), -1, dtype=torch.long)
                        for pos, idx in enumerate(bank_tensor.tolist()):
                            if 0 <= idx < pos_in_bank.numel():
                                pos_in_bank[idx] = pos
                        self_mask = pos_in_bank >= 0
                        if self_mask.any():
                            sim[self_mask, pos_in_bank[self_mask]] = -1e9
                        k = min(topk, int(bank_tensor.numel() - 1))
                        if k > 0:
                            topk_pos = torch.topk(sim, k=k, dim=1, largest=True).indices  # (N, k)
                            ref_indices = bank_tensor[topk_pos]  # (N, k)
                            print(f"[Transductive] retrieval_topk={k} (within bank)")

                    ref_indices = ref_indices.to(device=self.device)

                    # Pass-1: recompute MSM with restricted refs (bank / top-k in bank)
                    anomaly_maps_r_ref_list = []
                    for r in self.r_list:
                        anomaly_maps_l_ref = torch.tensor([]).double()
                        for l in Z_cache[int(r)].keys():
                            Z = Z_cache[int(r)][l].to(self.device)
                            print(f'layer-{l} mutual scoring (bank)...')
                            anomaly_maps_msm = MSM_with_refs(
                                Z=Z,
                                device=self.device,
                                topmin_min=self.transductive_msm_topmin_min,
                                topmin_max=self.transductive_msm_topmin_max,
                                ref_indices=ref_indices,
                                cdist_chunk_size=self.msm_cdist_chunk_size,
                            )
                            if fused_attn_all is not None:
                                anomaly_maps_msm = self.hier_fusion.rescore_patch_scores(
                                    anomaly_maps_msm,
                                    fused_attn_all,
                                    lambda_scale=lambda_all,
                                )
                            anomaly_maps_l_ref = torch.cat(
                                (anomaly_maps_l_ref, anomaly_maps_msm.unsqueeze(0).cpu()), dim=0
                            )
                            torch.cuda.empty_cache()
                        anomaly_maps_l_ref = torch.mean(anomaly_maps_l_ref, 0)
                        anomaly_maps_r_ref_list.append(anomaly_maps_l_ref)
                    
                    # Fuse with adaptive weights
                    anomaly_maps_r_ref = torch.stack(anomaly_maps_r_ref_list, dim=0)
                    anomaly_maps_patch = self._fuse_anomaly_maps_r(anomaly_maps_r_ref, f"{category}_transductive").to(self.device)
                    del anomaly_maps_r_ref, anomaly_maps_r_ref_list
                    torch.cuda.empty_cache()

            # Optional safe fusion / adaptive gating between Pass-0 and Pass-1.
            # This impacts the image-level score (ac_score) and, optionally, pixel maps.
            anomaly_maps_patch_pass0 = anomaly_maps_patch_base
            anomaly_maps_patch_pass1 = anomaly_maps_patch if (need_transductive and bank_indices is not None) else None

            fusion_w = 1.0
            if anomaly_maps_patch_pass1 is not None and self.transductive_safe_fusion:
                if self.transductive_adaptive_gate:
                    conf = float(transductive_knee_strength) * float(transductive_kept_ratio)
                    conf = max(0.0, min(1.0, conf))
                    thr = float(self.transductive_gate_threshold)
                    if conf < thr:
                        fusion_w = 0.0
                        print(
                            f"[Transductive] adaptive_gate: knee={float(transductive_knee_strength):.4f} "
                            f"kept={float(transductive_kept_ratio):.3f} conf={conf:.4f} < thr={thr:.4f} -> fallback Pass-0"
                        )
                    else:
                        fusion_w = conf
                        print(
                            f"[Transductive] adaptive_gate: knee={float(transductive_knee_strength):.4f} "
                            f"kept={float(transductive_kept_ratio):.3f} fusion_w={fusion_w:.4f}"
                        )
                else:
                    fusion_w = float(self.transductive_safe_fusion_weight)
                    fusion_w = max(0.0, min(1.0, fusion_w))
                    print(f"[Transductive] safe_fusion: fusion_w={fusion_w:.2f}")

            if anomaly_maps_patch_pass1 is None:
                anomaly_maps_patch_for_image = anomaly_maps_patch_pass0
            elif self.transductive_safe_fusion:
                anomaly_maps_patch_for_image = (1.0 - fusion_w) * anomaly_maps_patch_pass0 + fusion_w * anomaly_maps_patch_pass1
            else:
                anomaly_maps_patch_for_image = anomaly_maps_patch_pass1

            if (anomaly_maps_patch_pass1 is not None) and (not self.transductive_apply_to_pixel):
                anomaly_maps_patch_for_pixel = anomaly_maps_patch_pass0
                print("[Transductive] apply_to_pixel=False (pixel uses Pass-0 maps)")
            else:
                anomaly_maps_patch_for_pixel = anomaly_maps_patch_for_image

            # interpolate pixel maps (for pixel metrics / visualization)
            B, L = anomaly_maps_patch_for_pixel.shape
            H = int(np.sqrt(L))
            anomaly_maps_iter = F.interpolate(
                anomaly_maps_patch_for_pixel.view(B, 1, H, H),
                size=self.image_size,
                mode='bilinear',
                align_corners=True,
            )
            anomaly_maps = torch.cat((anomaly_maps, anomaly_maps_iter.cpu()), dim=0)

            # keep pass-0 maps for optional debugging/reporting
            if anomaly_maps_base is not None:
                B0, L0 = anomaly_maps_patch_pass0.shape
                H0 = int(np.sqrt(L0))
                maps0 = F.interpolate(
                    anomaly_maps_patch_pass0.view(B0, 1, H0, H0),
                    size=self.image_size,
                    mode='bilinear',
                    align_corners=True,
                )
                anomaly_maps_base = torch.cat((anomaly_maps_base, maps0.cpu()), dim=0)

            # store image-level pooling input (ac_score) consistent with chosen image route
            # NOTE: we only store the pooled score to avoid extra memory.
            if 'ac_scores_all' not in locals():
                ac_scores_all = []
            B1, L1 = anomaly_maps_patch_for_image.shape
            H1 = int(np.sqrt(L1))
            anomaly_maps_iter_img = F.interpolate(
                anomaly_maps_patch_for_image.view(B1, 1, H1, H1),
                size=self.image_size,
                mode='bilinear',
                align_corners=True,
            )
            ac_scores_all.append(anomaly_maps_iter_img.view(B1, -1).max(dim=1).values.detach().cpu())

            # free big tensors
            del patch_tokens_list
            if Z_cache is not None:
                del Z_cache
            torch.cuda.empty_cache()

        # save image features for optimizing classification
        # cls_save_path = os.path.join('./image_features/{}_{}.dat'.format(dataset, category))
        # with open(cls_save_path, 'wb') as f:
        #     pickle.dump([np.array(class_tokens)], f)
        end_time_all = time.time()
        print('MuSc: {}ms per image'.format((end_time_all-start_time_all)*1000/dataset_num))

        anomaly_maps = anomaly_maps.cpu().numpy()
        torch.cuda.empty_cache()

        # the number of unlabeled test images
        B = anomaly_maps.shape[0]

        # image-level anomaly score input for RsCIN
        if 'ac_scores_all' in locals() and ac_scores_all:
            ac_score = torch.cat(ac_scores_all, dim=0).numpy()
        else:
            ac_score = np.array(anomaly_maps).reshape(B, -1).max(-1)
        # RsCIN
        if self.dataset == 'visa':
            k_score = [1, 8, 9]
        elif self.dataset == 'mvtec_ad':
            k_score = [1, 2, 3]
        else:
            k_score = [1, 2, 3]
        scores_cls = RsCIN(ac_score, class_tokens, k_list=k_score)

        # Optionally dump image-level scores for RsCIN histogram visualization.
        # - before RsCIN: ac_score (pooled from anomaly maps)
        # - after RsCIN: scores_cls (output of RsCIN, before any further rescoring)
        if self.dump_rscin_hist:
            try:
                dump_root = os.path.join(self.output_dir, self.dump_rscin_hist_dir)
                os.makedirs(dump_root, exist_ok=True)
                np.savez_compressed(
                    os.path.join(dump_root, f'{category}.npz'),
                    dataset=str(self.dataset),
                    category=str(category),
                    backbone=str(self.model_name),
                    img_size=int(self.image_size),
                    k_score=np.array(k_score, dtype=np.int64),
                    image_paths=np.array(image_path_list, dtype=object),
                    gt_sp=np.array(gt_list, dtype=np.int64),
                    score_before_rscin=np.array(ac_score, dtype=np.float32),
                    score_after_rscin=np.array(scores_cls, dtype=np.float32),
                )
            except Exception as exc:
                print(f'[warn] failed to dump RsCIN hist scores for {category}: {exc}')

        # Optionally dump inputs for offline sweep of image-level rescoring.
        # This is intentionally saved BEFORE any image rescoring happens.
        if self.dump_rescore_inputs and cls_token_batches:
            try:
                dump_root = os.path.join(self.output_dir, self.dump_rescore_dir)
                os.makedirs(dump_root, exist_ok=True)
                cls_tensor = torch.cat(cls_token_batches, dim=0).cpu()
                torch.save(
                    {
                        'dataset': self.dataset,
                        'category': category,
                        'image_paths': image_path_list,
                        'gt_sp': np.array(gt_list, dtype=np.int64),
                        'scores_base': np.array(scores_cls, dtype=np.float32),
                        'cls_tokens': cls_tensor,
                        'cls_topk_ratio': float(self.hier_fusion.cls_topk_ratio) if self.hier_fusion is not None else None,
                    },
                    os.path.join(dump_root, f'{category}.pt'),
                )
            except Exception as exc:
                print(f'[warn] failed to dump rescore inputs for {category}: {exc}')

        if self.use_hier_fusion and self.hier_fusion is not None and cls_token_batches:
            cls_tensor = torch.cat(cls_token_batches, dim=0)
            scores_cls = self.hier_fusion.rescore_image_scores(scores_cls, cls_tensor)

        print('computing metrics...')
        pr_sp = np.array(scores_cls)
        gt_sp = np.array(gt_list)
        gt_px = torch.cat(img_masks, dim=0).numpy().astype(np.int32)
        pr_px = np.array(anomaly_maps)
        image_metric, pixel_metric = compute_metrics(gt_sp, pr_sp, gt_px, pr_px)
        auroc_sp, f1_sp, ap_sp = image_metric
        auroc_px, f1_px, ap_px, aupro = pixel_metric
        print(category)
        print('image-level, auroc:{}, f1:{}, ap:{}'.format(auroc_sp*100, f1_sp*100, ap_sp*100))
        print('pixel-level, auroc:{}, f1:{}, ap:{}, aupro:{}'.format(auroc_px*100, f1_px*100, ap_px*100, aupro*100))

        if self.vis:
            print('visualization...')
            self.visualization(image_path_list, gt_list, pr_px, category)
        if self.use_hier_fusion and self.hier_fusion is not None and self.hier_fusion_stats is not None:
            stats = self.hier_fusion.gather_stats()
            if stats is not None:
                self.hier_fusion_stats[category] = stats
                stats_dir = os.path.join(self.output_dir, 'stats')
                os.makedirs(stats_dir, exist_ok=True)
                stats_path = os.path.join(stats_dir, f'{category}_stats.pt')
                torch.save({k: (v.cpu() if isinstance(v, torch.Tensor) else v) for k, v in stats.items()}, stats_path)

        return image_metric, pixel_metric


    def main(self):
        auroc_sp_ls = []
        f1_sp_ls = []
        ap_sp_ls = []
        auroc_px_ls = []
        f1_px_ls = []
        ap_px_ls = []
        aupro_ls = []
        for category in self.categories:
            image_metric, pixel_metric = self.make_category_data(category=category,)
            auroc_sp, f1_sp, ap_sp = image_metric
            auroc_px, f1_px, ap_px, aupro = pixel_metric
            auroc_sp_ls.append(auroc_sp)
            f1_sp_ls.append(f1_sp)
            ap_sp_ls.append(ap_sp)
            auroc_px_ls.append(auroc_px)
            f1_px_ls.append(f1_px)
            ap_px_ls.append(ap_px)
            aupro_ls.append(aupro)
        # mean
        auroc_sp_mean = sum(auroc_sp_ls) / len(auroc_sp_ls)
        f1_sp_mean = sum(f1_sp_ls) / len(f1_sp_ls)
        ap_sp_mean = sum(ap_sp_ls) / len(ap_sp_ls)
        auroc_px_mean = sum(auroc_px_ls) / len(auroc_px_ls)
        f1_px_mean = sum(f1_px_ls) / len(f1_px_ls)
        ap_px_mean = sum(ap_px_ls) / len(ap_px_ls)
        aupro_mean = sum(aupro_ls) / len(aupro_ls)

        for i, category in enumerate(self.categories):
            print(category)
            print('image-level, auroc:{}, f1:{}, ap:{}'.format(auroc_sp_ls[i]*100, f1_sp_ls[i]*100, ap_sp_ls[i]*100))
            print('pixel-level, auroc:{}, f1:{}, ap:{}, aupro:{}'.format(auroc_px_ls[i]*100, f1_px_ls[i]*100, ap_px_ls[i]*100, aupro_ls[i]*100))
        print('mean')
        print('image-level, auroc:{}, f1:{}, ap:{}'.format(auroc_sp_mean*100, f1_sp_mean*100, ap_sp_mean*100))
        print('pixel-level, auroc:{}, f1:{}, ap:{}, aupro:{}'.format(auroc_px_mean*100, f1_px_mean*100, ap_px_mean*100, aupro_mean*100))
        
        # save in excel
        if self.save_excel:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "MuSc_results"
            sheet.cell(row=1,column=2,value='auroc_px')
            sheet.cell(row=1,column=3,value='f1_px')
            sheet.cell(row=1,column=4,value='ap_px')
            sheet.cell(row=1,column=5,value='aupro')
            sheet.cell(row=1,column=6,value='auroc_sp')
            sheet.cell(row=1,column=7,value='f1_sp')
            sheet.cell(row=1,column=8,value='ap_sp')
            for col_index in range(2):
                for row_index in range(len(self.categories)):
                    if col_index == 0:
                        sheet.cell(row=row_index+2,column=col_index+1,value=self.categories[row_index])
                    else:
                        sheet.cell(row=row_index+2,column=col_index+1,value=auroc_px_ls[row_index]*100)
                        sheet.cell(row=row_index+2,column=col_index+2,value=f1_px_ls[row_index]*100)
                        sheet.cell(row=row_index+2,column=col_index+3,value=ap_px_ls[row_index]*100)
                        sheet.cell(row=row_index+2,column=col_index+4,value=aupro_ls[row_index]*100)
                        sheet.cell(row=row_index+2,column=col_index+5,value=auroc_sp_ls[row_index]*100)
                        sheet.cell(row=row_index+2,column=col_index+6,value=f1_sp_ls[row_index]*100)
                        sheet.cell(row=row_index+2,column=col_index+7,value=ap_sp_ls[row_index]*100)
                    if row_index == len(self.categories)-1:
                        if col_index == 0:
                            sheet.cell(row=row_index+3,column=col_index+1,value='mean')
                        else:
                            sheet.cell(row=row_index+3,column=col_index+1,value=auroc_px_mean*100)
                            sheet.cell(row=row_index+3,column=col_index+2,value=f1_px_mean*100)
                            sheet.cell(row=row_index+3,column=col_index+3,value=ap_px_mean*100)
                            sheet.cell(row=row_index+3,column=col_index+4,value=aupro_mean*100)
                            sheet.cell(row=row_index+3,column=col_index+5,value=auroc_sp_mean*100)
                            sheet.cell(row=row_index+3,column=col_index+6,value=f1_sp_mean*100)
                            sheet.cell(row=row_index+3,column=col_index+7,value=ap_sp_mean*100)
            # 重跑或清理目录后确保结果路径存在，避免保存 Excel 时报 FileNotFoundError。
            os.makedirs(self.output_dir, exist_ok=True)
            workbook.save(os.path.join(self.output_dir, 'results.xlsx'))

        # Save AMSF statistics if enabled
        if self.adaptive_r_fusion_stats:
            stats_dir = os.path.join(self.output_dir, 'amsf_stats')
            os.makedirs(stats_dir, exist_ok=True)
            import json
            with open(os.path.join(stats_dir, 'weight_stats.json'), 'w') as f:
                json.dump(self.adaptive_r_fusion_stats, f, indent=2)

        # Save ALF statistics if enabled
        if self.adaptive_layer_fusion_stats:
            stats_dir = os.path.join(self.output_dir, 'alf_stats')
            os.makedirs(stats_dir, exist_ok=True)
            import json
            with open(os.path.join(stats_dir, 'weight_stats.json'), 'w') as f:
                json.dump(self.adaptive_layer_fusion_stats, f, indent=2)

        if self.use_hier_fusion and self.hier_fusion_stats:
            stats_dir = os.path.join(self.output_dir, 'stats')
            os.makedirs(stats_dir, exist_ok=True)
            summary_path = os.path.join(stats_dir, 'summary.pt')
            torch.save(
                {
                    category: {k: (v.cpu() if isinstance(v, torch.Tensor) else v) for k, v in stats.items()}
                    for category, stats in self.hier_fusion_stats.items()
                },
                summary_path,
            )
